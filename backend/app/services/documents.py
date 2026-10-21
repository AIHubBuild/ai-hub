"""Document services: text extraction and DOCX generation.

Extraction ports the legacy ATS `file_utils.py` behavior
(`docs/09_Migration_Analysis.md` §2.1.8, §3.2 KEEP) onto FastAPI-shaped
inputs: bytes + original filename. DOCX generation ports the
`generate_documents.py` section-based styling (§3.2 REFACTOR) with the
hardcoded personal-data heuristics removed — no location guessing, no
name assumptions.

Validation rules (legacy-preserved): extension whitelist (.pdf, .docx,
.txt), a configured size limit (default 5 MB), and non-empty extracted
text.
"""

from io import BytesIO

import re

from docx import Document as DocxDocument
from docx.shared import Pt
from pypdf import PdfReader

from app.services.errors import (
    DocumentError,
    EmptyDocumentError,
    InvalidFileTypeError,
    UnreadableDocumentError,
)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

DEFAULT_MAX_SIZE_BYTES = 5 * 1024 * 1024

_MARKDOWN_HEADING = re.compile(r"^\s*#{1,6}\s*")
_MARKDOWN_EMPHASIS = re.compile(r"\*\*|__")
_MARKDOWN_STAR = re.compile(r"(?<!\*)\*(?!\*)")
_DECORATIVE_LINE = re.compile(r"^\s*(?:---+|___+|\*\*\*+)\s*$")
_BULLET_LINE = re.compile(r"^\s*[-*•]\s+")

# Section headings recognized when rendering a plain-text tailored
# resume into DOCX. Compared case-insensitively after light cleanup.
SECTION_HEADINGS = {
    "summary",
    "professional summary",
    "skills",
    "technical skills",
    "core competencies",
    "experience",
    "work experience",
    "professional experience",
    "projects",
    "education",
    "certifications",
    "certificates",
    "languages",
}


def extract_text(filename: str, data: bytes, max_size_bytes: int = DEFAULT_MAX_SIZE_BYTES) -> str:
    """Extract plain text from an uploaded resume file.

    Raises InvalidFileTypeError for unsupported extensions and
    DocumentError subclasses for size/content problems.
    """
    if len(data) > max_size_bytes:
        raise DocumentError(
            f"File exceeds the maximum allowed size of "
            f"{max_size_bytes // (1024 * 1024)} MB."
        )

    extension = _extension_of(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise InvalidFileTypeError(
            "Unsupported file type. Supported formats: PDF, DOCX, TXT."
        )

    if extension == ".txt":
        text = _extract_txt(data)
    elif extension == ".pdf":
        text = _extract_pdf(data)
    else:
        text = _extract_docx(data)

    if not text.strip():
        raise EmptyDocumentError("No readable text found in the uploaded file.")

    return text.strip()


def _clean_generated_line(line: str) -> str:
    """Remove Markdown/decorative markers without changing facts."""
    cleaned = _MARKDOWN_HEADING.sub("", line.strip())
    cleaned = _MARKDOWN_EMPHASIS.sub("", cleaned)
    cleaned = _MARKDOWN_STAR.sub("", cleaned)
    return cleaned.strip()


def _is_decorative_line(line: str) -> bool:
    return _DECORATIVE_LINE.match(line.strip()) is not None


def build_tailored_resume_docx(text: str) -> bytes:
    """Render tailored resume text with clean, conventional DOCX formatting."""
    document = DocxDocument()

    first = True
    for raw_line in text.splitlines():
        if not raw_line.strip() or _is_decorative_line(raw_line):
            continue
        is_bullet = _BULLET_LINE.match(raw_line) is not None
        line = _clean_generated_line(raw_line)
        if not line:
            continue
        if _is_contact_or_date_line(line):
            _append_paragraph(document, line, bold=False, size_pt=11)
            continue
        if first:
            first = False
            _append_paragraph(document, line, bold=True, size_pt=20)
            continue
        if _is_section_heading(line):
            _append_paragraph(document, line, bold=True, size_pt=14)
            continue
        if is_bullet:
            line = _BULLET_LINE.sub("", line).strip()
            if line:
                document.add_paragraph(line, style="List Bullet")
            continue
        _append_paragraph(document, line, bold=False, size_pt=11)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def validate_docx(content: bytes) -> bool:
    """True when `content` parses as a DOCX with readable text."""
    try:
        document = DocxDocument(BytesIO(content))
    except Exception:
        return False
    return any(paragraph.text.strip() for paragraph in document.paragraphs)


def build_cover_letter_docx(text: str) -> bytes:
    """Render cover-letter text as conventional professional paragraphs."""
    document = DocxDocument()

    for raw_line in text.splitlines():
        if _is_decorative_line(raw_line):
            continue
        line = _clean_generated_line(raw_line)
        if line:
            _append_paragraph(document, line, bold=False, size_pt=11)
        else:
            document.add_paragraph()

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _is_contact_or_date_line(line: str) -> bool:
    if "@" in line or "linkedin.com" in line.lower():
        return True
    if re.search(r"\+?[\d\s().\-]{9,}", line):
        return True
    # Date ranges like "2020 - 2023" or "June 2020 – Present".
    if re.search(r"(19|20)\d{2}\s*[-–—]\s*((19|20)\d{2}|present)", line, re.IGNORECASE):
        return True
    return False


def _extension_of(filename: str) -> str:
    return ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnreadableDocumentError("Could not decode the uploaded text file.")


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise UnreadableDocumentError("Could not read the uploaded PDF.") from exc
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            raise UnreadableDocumentError(
                "Could not extract text from the uploaded PDF."
            ) from exc
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    try:
        document = DocxDocument(BytesIO(data))
    except Exception as exc:
        raise UnreadableDocumentError("Could not read the uploaded DOCX.") from exc
    return "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )


def _is_section_heading(line: str) -> bool:
    normalized = line.rstrip(":").strip().lower()
    return normalized in SECTION_HEADINGS


def _append_paragraph(document: DocxDocument, text: str, *, bold: bool, size_pt: int) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
