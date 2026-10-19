"""End-to-end tests for the resume_ats API.

The OpenAI provider is replaced with a scripted test double through the
same ai_core wiring the application uses — no network calls. The
deterministic engine (not the double) owns the ATS score.

All endpoints are protected by Supabase-style HS256 JWTs validated
against the app-bound `supabase_jwt_secret`; tokens here are minted
with PyJWT exactly as Supabase would sign them. Cross-user isolation
is exercised throughout with two distinct user identities.
"""

import time

import jwt
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.ai_core import (
    AICore,
    DomainModelConfig,
    create_ai_core,
)
from app.ai_core.providers.openai import ProviderError
from app.ai_core.run_records_db import SQLAlchemyRunRecordStore
from app.core.config import Settings
from app.domains.resume_ats.prompts import register_prompts
from app.domains.resume_ats.workflow import ATS_WORKFLOW_ID, build_ats_workflow
from app.main import create_app
from app.models import Base
from app.models.admin import UserAccountStatus
from app.models.resume import AtsAnalysisRecord, Resume

RESUME_TEXT = (
    "Jane Doe\njane@example.com\n"
    "Experienced with Active Directory, Microsoft 365, VPN and "
    "network troubleshooting."
)
JOB_DESCRIPTION = (
    "We need Active Directory, Microsoft 365, VPN, PowerShell and "
    "Intune experience."
)

FAKE_RESPONSES = [
    "Add PowerShell automation projects and lab exposure.",
    "Jane Doe\nTailored resume content.",
    "Tailored cover letter content.\nJane Doe",
]

TEST_SECRET = "test-jwt-secret-0123456789abcdef0123456789abcdef0123456789abcdef"

USER_A = "550e8400-e29b-41d4-a716-446655440000"
USER_B = "660e8400-e29b-41d4-a716-446655440001"


class FakeProvider:
    def __init__(self, responses=None, error=None):
        self.responses = list(responses or [])
        self.error = error
        self.calls = 0

    def complete(self, *, system, user, model, temperature):
        self.calls += 1
        if self.error is not None:
            raise self.error
        if not self.responses:
            raise AssertionError("No canned response available for LLM call")
        return self.responses.pop(0)


def make_token(user_id: str = USER_A, *, email: str = "alex@example.com") -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "exp": int(time.time()) + 3600,
    }
    return jwt.encode(payload, TEST_SECRET, algorithm="HS256")


def auth_headers(user_id: str = USER_A) -> dict[str, str]:
    return {"Authorization": f"Bearer {make_token(user_id)}"}


def make_settings(tmp_path, **overrides) -> Settings:
    values = dict(
        database_url=f"sqlite:///{tmp_path / 'test.db'}",
        openai_api_key="test-key",
        supabase_jwt_secret=TEST_SECRET,
    )
    values.update(overrides)
    return Settings(**values)


def build_app(tmp_path, settings: Settings) -> FastAPI:
    app = create_app(settings)
    Base.metadata.create_all(app.state.engine)
    return app


def install_fake_provider(app: FastAPI, provider: FakeProvider) -> None:
    """Rebuild app.state.ai_core around a scripted provider double."""
    run_store = SQLAlchemyRunRecordStore(app.state.session_factory)
    ai_core: AICore = create_ai_core(run_store=run_store)
    ai_core.llm_factory.register_config(
        DomainModelConfig(domain="resume_ats", provider="openai", model="gpt-4.1-mini")
    )
    ai_core.llm_factory.register_provider("openai", provider)
    register_prompts(ai_core.prompts)
    client = ai_core.llm_factory.get_client("resume_ats")
    graph = build_ats_workflow(client, ai_core.prompts)
    ai_core.workflows.register(ATS_WORKFLOW_ID, graph.invoke)
    app.state.ai_core = ai_core


@pytest.fixture
def client(tmp_path):
    app = build_app(tmp_path, make_settings(tmp_path))
    install_fake_provider(app, FakeProvider(responses=list(FAKE_RESPONSES)))
    with TestClient(app) as test_client:
        yield test_client, app


def upload_resume(
    client: TestClient,
    content: str = RESUME_TEXT,
    filename="resume.txt",
    user_id: str = USER_A,
):
    return client.post(
        "/api/v1/resumes",
        files={"file": (filename, content.encode("utf-8"), "text/plain")},
        headers=auth_headers(user_id),
    )


# ---------------------------------------------------------------
# Authentication requirements
# ---------------------------------------------------------------


def test_upload_requires_authentication(client):
    test_client, _ = client
    response = test_client.post(
        "/api/v1/resumes",
        files={"file": ("resume.txt", RESUME_TEXT.encode("utf-8"), "text/plain")},
    )

    assert response.status_code == 401
    assert response.json()["success"] is False


def test_resume_detail_requires_authentication(client):
    test_client, _ = client
    response = test_client.get("/api/v1/resumes/550e8400-e29b-41d4-a716-446655440000")

    assert response.status_code == 401


def test_analysis_requires_authentication(client):
    test_client, _ = client
    response = test_client.post(
        "/api/v1/resumes/550e8400-e29b-41d4-a716-446655440000/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
    )

    assert response.status_code == 401


def test_download_requires_authentication(client):
    test_client, _ = client
    response = test_client.get(
        "/api/v1/analyses/550e8400-e29b-41d4-a716-446655440000/tailored-resume.docx"
    )

    assert response.status_code == 401


def test_suspended_user_cannot_access_resume_ats(client):
    test_client, app = client
    with app.state.session_factory() as session:
        session.add(UserAccountStatus(user_id=USER_A, status="suspended"))
        session.commit()

    response = upload_resume(test_client, user_id=USER_A)

    assert response.status_code == 403
    assert response.json() == {
        "success": False,
        "message": "This account is suspended.",
    }


def test_invalid_token_is_rejected(client):
    test_client, _ = client
    response = test_client.get(
        "/api/v1/resumes/550e8400-e29b-41d4-a716-446655440000",
        headers={"Authorization": "Bearer not-a-jwt"},
    )

    assert response.status_code == 401


# ---------------------------------------------------------------
# Upload
# ---------------------------------------------------------------


def test_upload_txt_resume(client):
    test_client, _ = client
    response = upload_resume(test_client)

    assert response.status_code == 201
    body = response.json()
    assert body["success"] is True
    assert body["data"]["file_name"] == "resume.txt"
    assert body["data"]["resume_id"]
    assert body["data"]["extracted_characters"] > 0


def test_upload_persists_authenticated_user_id(client):
    test_client, app = client
    response = upload_resume(test_client, user_id=USER_B)
    resume_id = response.json()["data"]["resume_id"]

    with app.state.session_factory() as session:
        resume = session.get(Resume, resume_id)

    assert resume.user_id == USER_B


def test_upload_rejects_unsupported_extension(client):
    test_client, _ = client
    response = upload_resume(test_client, filename="resume.exe")

    assert response.status_code == 415
    assert response.json()["success"] is False


def test_upload_rejects_missing_extension(client):
    test_client, _ = client
    response = upload_resume(test_client, filename="resume")

    assert response.status_code == 415


def test_upload_rejects_oversized_file(tmp_path):
    settings = make_settings(tmp_path, max_upload_size_mb=1)
    app = build_app(tmp_path, settings)
    with TestClient(app) as test_client:
        response = upload_resume(
            test_client, content="x" * (1024 * 1024 + 1)
        )

    assert response.status_code == 413


def test_upload_rejects_empty_file(client):
    test_client, _ = client
    response = upload_resume(test_client, content="   \n ")

    assert response.status_code == 422


# ---------------------------------------------------------------
# Resume detail
# ---------------------------------------------------------------


def test_get_resume_returns_metadata_and_no_analysis(client):
    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]

    response = test_client.get(
        f"/api/v1/resumes/{resume_id}", headers=auth_headers()
    )

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["resume_id"] == resume_id
    assert data["file_name"] == "resume.txt"
    assert data["latest_analysis"] is None


def test_get_resume_unknown_returns_404(client):
    test_client, _ = client
    response = test_client.get(
        "/api/v1/resumes/550e8400-e29b-41d4-a716-446655440000",
        headers=auth_headers(),
    )

    assert response.status_code == 404
    assert response.json()["success"] is False


def test_get_resume_malformed_id_returns_404_not_500(client):
    test_client, _ = client
    response = test_client.get(
        "/api/v1/resumes/not-a-uuid", headers=auth_headers()
    )

    assert response.status_code == 404
    assert response.json()["success"] is False


def test_get_resume_of_other_user_returns_404(client):
    test_client, _ = client
    resume_id = upload_resume(test_client, user_id=USER_A).json()["data"]["resume_id"]

    response = test_client.get(
        f"/api/v1/resumes/{resume_id}", headers=auth_headers(USER_B)
    )

    assert response.status_code == 404
    assert response.json()["success"] is False


# ---------------------------------------------------------------
# ATS analysis
# ---------------------------------------------------------------


def test_analysis_runs_workflow_and_returns_six_sections(client):
    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(),
    )

    assert response.status_code == 201
    body = response.json()
    assert body["success"] is True
    data = body["data"]
    # Three of the five terms in the experience requirement match; the
    # requirement is partial, and evidence confidence is limited because
    # the resume does not show dated experience.
    assert data["ats_score"] == 39
    assert set(data["matching_keywords"]) == {
        "active directory",
        "microsoft 365",
        "vpn",
    }
    assert set(data["missing_keywords"]) == {"powershell", "intune"}
    assert data["improvement_suggestions"] == FAKE_RESPONSES[0]
    assert data["tailored_resume"] == FAKE_RESPONSES[1]
    assert data["cover_letter"] == FAKE_RESPONSES[2]
    assert data["run_id"]


def test_analysis_persists_run_record_for_authenticated_user(client):
    test_client, app = client
    resume_id = upload_resume(test_client, user_id=USER_A).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(USER_A),
    )
    run_id = response.json()["data"]["run_id"]

    record = app.state.ai_core.run_records.get(run_id)
    assert record is not None
    assert record.workflow_id == ATS_WORKFLOW_ID
    assert record.domain == "resume_ats"
    assert record.user_id == USER_A
    assert record.status.value == "success"
    assert record.input_state["resume_text"] == RESUME_TEXT


def test_analysis_unknown_resume_returns_404(client):
    test_client, _ = client
    response = test_client.post(
        "/api/v1/resumes/550e8400-e29b-41d4-a716-446655440000/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(),
    )

    assert response.status_code == 404


def test_analysis_of_other_users_resume_returns_404(client):
    test_client, app = client
    resume_id = upload_resume(test_client, user_id=USER_A).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(USER_B),
    )

    assert response.status_code == 404
    with app.state.session_factory() as session:
        assert session.query(AtsAnalysisRecord).count() == 0


def test_analysis_rejects_empty_job_description(client):
    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": "   "},
        headers=auth_headers(),
    )

    assert response.status_code == 422


def test_analysis_provider_failure_returns_502_and_error_run_record(tmp_path):
    settings = make_settings(tmp_path)
    app = build_app(tmp_path, settings)
    install_fake_provider(
        app, FakeProvider(error=ProviderError("AI provider returned HTTP 500.", status=500))
    )
    with TestClient(app) as test_client:
        resume_id = upload_resume(test_client).json()["data"]["resume_id"]
        response = test_client.post(
            f"/api/v1/resumes/{resume_id}/ats-analysis",
            json={"job_description": JOB_DESCRIPTION},
            headers=auth_headers(),
        )

    assert response.status_code == 502
    assert response.json()["success"] is False


def test_analysis_without_provider_returns_503(tmp_path):
    settings = make_settings(tmp_path, openai_api_key="")
    app = build_app(tmp_path, settings)
    with TestClient(app) as test_client:
        resume_id = upload_resume(test_client).json()["data"]["resume_id"]
        response = test_client.post(
            f"/api/v1/resumes/{resume_id}/ats-analysis",
            json={"job_description": JOB_DESCRIPTION},
            headers=auth_headers(),
        )

    assert response.status_code == 503
    message = response.json()["message"]
    assert "AI_API_KEY" in message
    assert "DEEPSEEK_API_KEY" in message


def test_no_provider_means_no_fabricated_analysis_or_run_record(tmp_path):
    """A 503 response must not create analysis rows or run records."""
    from app.models.resume import AtsAnalysisRecord

    settings = make_settings(tmp_path, openai_api_key="")
    app = build_app(tmp_path, settings)
    with TestClient(app) as test_client:
        resume_id = upload_resume(test_client).json()["data"]["resume_id"]
        response = test_client.post(
            f"/api/v1/resumes/{resume_id}/ats-analysis",
            json={"job_description": JOB_DESCRIPTION},
            headers=auth_headers(),
        )
        assert response.status_code == 503

    with app.state.session_factory() as session:
        assert session.query(AtsAnalysisRecord).count() == 0
    assert app.state.ai_core.run_records.list_for_user(USER_A) == []


def test_provider_failure_creates_no_analysis_row_but_error_run_record(tmp_path):
    """A provider failure returns 502: no analysis data, error record kept."""
    from app.ai_core.run_records import RunStatus
    from app.models.resume import AtsAnalysisRecord

    settings = make_settings(tmp_path)
    app = build_app(tmp_path, settings)
    install_fake_provider(
        app, FakeProvider(error=ProviderError("AI provider returned HTTP 500.", status=500))
    )
    with TestClient(app) as test_client:
        resume_id = upload_resume(test_client).json()["data"]["resume_id"]
        response = test_client.post(
            f"/api/v1/resumes/{resume_id}/ats-analysis",
            json={"job_description": JOB_DESCRIPTION},
            headers=auth_headers(),
        )
        assert response.status_code == 502
        assert response.json()["success"] is False

    with app.state.session_factory() as session:
        assert session.query(AtsAnalysisRecord).count() == 0
    records = app.state.ai_core.run_records.list_for_user(USER_A)
    assert len(records) == 1
    assert records[0].status is RunStatus.ERROR


def test_analysis_response_includes_explainability_fields(client):
    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(),
    )

    data = response.json()["data"]
    assert data["score_label"] == "Estimated resume-to-job match"
    assert data["partial_keywords"] == []
    assert data["job_description_warning"] is None
    assert data["requirement_coverage"] == 50
    assert data["evidence_confidence"] < 100
    requirements = data["requirement_analysis"]
    experience = next(
        r for r in requirements if "Active Directory" in r["term"]
    )
    assert experience["status"] == "partial"
    assert experience["evidence"]
    assert "powershell" in experience["detail"]
    assert "powershell" in data["missing_keywords"]
    assert data["score_breakdown"]


def test_analysis_response_includes_writing_review(client):
    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(),
    )

    data = response.json()["data"]
    review = data["writing_review"]
    assert review is not None
    assert review["passed"] is True
    assert review["issues"] == []


def test_analysis_accepts_style_notes(client):
    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={
            "job_description": JOB_DESCRIPTION,
            "style_notes": "Prefer UK spelling.",
        },
        headers=auth_headers(),
    )

    assert response.status_code == 201
    assert response.json()["data"]["analysis_id"]


def test_boilerplate_job_description_warns_and_scores_anyway(client):
    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]
    boilerplate_jd = (
        "Apply now\n"
        "We use cookies to improve your experience\n"
        "Privacy policy\n"
        "Active Directory experience required."
    )

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": boilerplate_jd},
        headers=auth_headers(),
    )

    assert response.status_code == 201
    data = response.json()["data"]
    assert data["job_description_warning"] is not None
    assert "boilerplate" in data["job_description_warning"]
    # The genuine requirement still drives the analysis, but a single
    # matched term does not produce a perfect score.
    assert data["ats_score"] == 79
    assert data["matching_keywords"] == ["active directory"]


def test_synonym_match_is_reported_as_partial(client):
    test_client, _ = client
    resume_id = upload_resume(
        test_client, content="Provided help desk support for office users."
    ).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": "Technical support experience required."},
        headers=auth_headers(),
    )

    data = response.json()["data"]
    assert data["partial_keywords"] == ["technical support"]
    requirement = next(
        r for r in data["requirement_analysis"] if "Technical support" in r["term"]
    )
    assert requirement["status"] == "partial"
    assert requirement["evidence"]


def test_genuine_zero_score_reported_with_missing_keywords_and_breakdown(client):
    test_client, _ = client
    resume_id = upload_resume(
        test_client, content="Worked in retail customer service."
    ).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={
            "job_description": (
                "We need Microsoft 365, VPN, Teams and Windows experience."
            )
        },
        headers=auth_headers(),
    )

    assert response.status_code == 201
    data = response.json()["data"]
    # A resume that contains none of the detected requirements scores 0.
    assert data["ats_score"] == 0
    assert data["score_label"] == "Estimated resume-to-job match"
    assert data["matching_keywords"] == []
    assert data["missing_keywords"]
    assert data["score_breakdown"]
    assert all(0 <= entry["score"] <= 100 for entry in data["score_breakdown"])
    assert all(entry["total"] > 0 for entry in data["score_breakdown"])


def test_non_zero_score_reported_when_resume_covers_requirements(client):
    test_client, _ = client
    resume_id = upload_resume(
        test_client,
        content=(
            "Managed Microsoft 365 tenants, VPN gateways, Microsoft Teams "
            "deployments, and Windows 11 fleets."
        ),
    ).json()["data"]["resume_id"]

    response = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={
            "job_description": (
                "We need Microsoft 365, VPN, Teams and Windows experience."
            )
        },
        headers=auth_headers(),
    )

    assert response.status_code == 201
    data = response.json()["data"]
    # All four terms match, but the score is not 100 because the resume
    # has no dated experience, qualification, or achievement evidence.
    assert data["ats_score"] == 81
    assert 0 < data["ats_score"] < 100
    assert data["matching_keywords"]
    assert data["missing_keywords"] == []
    assert data["evidence_confidence"] < 100


def test_upload_docx_through_api(client):
    from io import BytesIO

    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("Resume with Active Directory experience.")
    buffer = BytesIO()
    document.save(buffer)

    test_client, _ = client
    response = test_client.post(
        "/api/v1/resumes",
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument"
                ".wordprocessingml.document",
            )
        },
        headers=auth_headers(),
    )

    assert response.status_code == 201
    data = response.json()["data"]
    assert data["file_name"] == "resume.docx"
    assert data["extracted_characters"] > 0


def test_wiring_registers_workflow_only_with_provider(tmp_path):
    app = build_app(tmp_path, make_settings(tmp_path))
    assert ATS_WORKFLOW_ID in app.state.ai_core.workflows._workflows

    app_without_key = build_app(tmp_path, make_settings(tmp_path, openai_api_key=""))
    assert ATS_WORKFLOW_ID not in app_without_key.state.ai_core.workflows._workflows


# ---------------------------------------------------------------
# DOCX download
# ---------------------------------------------------------------


def test_download_tailored_resume_docx(client):
    from io import BytesIO

    from docx import Document as DocxDocument

    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]
    analysis_id = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(),
    ).json()["data"]["analysis_id"]

    response = test_client.get(
        f"/api/v1/analyses/{analysis_id}/tailored-resume.docx",
        headers=auth_headers(),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    assert "attachment" in response.headers["content-disposition"]
    assert response.content.startswith(b"PK")

    parsed = DocxDocument(BytesIO(response.content))
    paragraphs = [p.text for p in parsed.paragraphs if p.text.strip()]
    assert "Tailored resume content." in paragraphs


def test_download_unknown_analysis_returns_404(client):
    test_client, _ = client
    response = test_client.get(
        "/api/v1/analyses/550e8400-e29b-41d4-a716-446655440000/tailored-resume.docx",
        headers=auth_headers(),
    )

    assert response.status_code == 404


def test_download_other_users_analysis_returns_404(client):
    test_client, _ = client
    resume_id = upload_resume(test_client, user_id=USER_A).json()["data"]["resume_id"]
    analysis_id = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(USER_A),
    ).json()["data"]["analysis_id"]

    response = test_client.get(
        f"/api/v1/analyses/{analysis_id}/tailored-resume.docx",
        headers=auth_headers(USER_B),
    )

    assert response.status_code == 404


def test_download_analysis_without_tailored_resume_returns_409(client):
    test_client, app = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]

    with app.state.session_factory() as session:
        record = AtsAnalysisRecord(
            resume_id=resume_id,
            run_id="550e8400-e29b-41d4-a716-446655440001",
            job_description=JOB_DESCRIPTION,
            ats_score=50,
            matching_keywords=[],
            missing_keywords=[],
            tailored_resume=None,
        )
        session.add(record)
        session.commit()
        analysis_id = record.id

    response = test_client.get(
        f"/api/v1/analyses/{analysis_id}/tailored-resume.docx",
        headers=auth_headers(),
    )

    assert response.status_code == 409


def test_download_malformed_analysis_id_returns_404_not_500(client):
    test_client, _ = client
    response = test_client.get(
        "/api/v1/analyses/not-a-uuid/tailored-resume.docx",
        headers=auth_headers(),
    )

    assert response.status_code == 404


def test_download_cover_letter_docx(client):
    from io import BytesIO

    from docx import Document as DocxDocument

    test_client, _ = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]
    analysis_id = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(),
    ).json()["data"]["analysis_id"]

    response = test_client.get(
        f"/api/v1/analyses/{analysis_id}/tailored-cover-letter.docx",
        headers=auth_headers(),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    assert "tailored-cover-letter.docx" in response.headers["content-disposition"]
    assert "attachment" in response.headers["content-disposition"]
    assert response.content.startswith(b"PK")

    parsed = DocxDocument(BytesIO(response.content))
    paragraphs = [p.text for p in parsed.paragraphs if p.text.strip()]
    assert "Tailored cover letter content." in paragraphs


def test_download_cover_letter_does_not_regenerate_or_reanalyze(client):
    """Downloading must reuse stored content without new AI calls."""
    test_client, app = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]
    analysis_id = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(),
    ).json()["data"]["analysis_id"]

    runs_before = len(
        app.state.ai_core.run_records.list_for_user(USER_A)
    )
    response = test_client.get(
        f"/api/v1/analyses/{analysis_id}/tailored-cover-letter.docx",
        headers=auth_headers(),
    )
    runs_after = len(app.state.ai_core.run_records.list_for_user(USER_A))

    assert response.status_code == 200
    assert runs_after == runs_before


def test_download_cover_letter_unknown_analysis_returns_404(client):
    test_client, _ = client
    response = test_client.get(
        "/api/v1/analyses/550e8400-e29b-41d4-a716-446655440000/tailored-cover-letter.docx",
        headers=auth_headers(),
    )

    assert response.status_code == 404


def test_download_cover_letter_requires_authentication(client):
    test_client, _ = client
    response = test_client.get(
        "/api/v1/analyses/550e8400-e29b-41d4-a716-446655440000/tailored-cover-letter.docx"
    )

    assert response.status_code == 401


def test_download_cover_letter_without_content_returns_409(client):
    test_client, app = client
    resume_id = upload_resume(test_client).json()["data"]["resume_id"]

    with app.state.session_factory() as session:
        record = AtsAnalysisRecord(
            resume_id=resume_id,
            run_id="550e8400-e29b-41d4-a716-446655440001",
            job_description=JOB_DESCRIPTION,
            ats_score=50,
            matching_keywords=[],
            missing_keywords=[],
            cover_letter=None,
        )
        session.add(record)
        session.commit()
        analysis_id = record.id

    response = test_client.get(
        f"/api/v1/analyses/{analysis_id}/tailored-cover-letter.docx",
        headers=auth_headers(),
    )

    assert response.status_code == 409


def test_download_cover_letter_other_user_returns_404(client):
    test_client, _ = client
    resume_id = upload_resume(test_client, user_id=USER_A).json()["data"]["resume_id"]
    analysis_id = test_client.post(
        f"/api/v1/resumes/{resume_id}/ats-analysis",
        json={"job_description": JOB_DESCRIPTION},
        headers=auth_headers(USER_A),
    ).json()["data"]["analysis_id"]

    response = test_client.get(
        f"/api/v1/analyses/{analysis_id}/tailored-cover-letter.docx",
        headers=auth_headers(USER_B),
    )

    assert response.status_code == 404
