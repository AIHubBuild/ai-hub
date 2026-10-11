"""Tests for the shared documents service (extraction + DOCX)."""

from io import BytesIO

import pytest
from docx import Document as DocxDocument

from app.services import documents
from app.services.errors import (
    DocumentError,
    EmptyDocumentError,
    InvalidFileTypeError,
    UnreadableDocumentError,
)

RESUME_TEXT = (
    "John Doe\n"
    "john.doe@example.com | +49 123 456\n"
    "Professional Summary:\n"
    "IT support specialist with Active Directory experience.\n"
    "Skills:\n"
    "- Windows administration\n"
    "- VPN troubleshooting\n"
    "Experience:\n"
    "Service desk analyst at Example GmbH."
)


def make_docx(text: str) -> bytes:
    document = DocxDocument()
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf() -> bytes:
    # Minimal single-page PDF with extractable text.
    content = b"Hello PDF resume content with Active Directory."
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\nBT /F1 12 Tf 72 720 Td (%s) Tj ET\nendstream"
        % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (index, obj)
    xref = len(pdf)
    pdf += b"xref\n0 %d\n" % (len(objects) + 1)
    pdf += b"0000000000 65535 f \n"
    for offset in offsets:
        pdf += b"%010d 00000 n \n" % offset
    pdf += (
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, xref)
    )
    return bytes(pdf)


def test_extract_txt_returns_trimmed_text():
    text = documents.extract_text("resume.txt", RESUME_TEXT.encode("utf-8"))
    assert text.startswith("John Doe")
    assert "Active Directory" in text


def test_extract_docx_returns_paragraph_text():
    data = make_docx(RESUME_TEXT)
    text = documents.extract_text("resume.docx", data)
    assert "John Doe" in text
    assert "Windows administration" in text


def test_extract_pdf_returns_text():
    text = documents.extract_text("resume.pdf", make_pdf())
    assert "Active Directory" in text


def test_extract_text_rejects_unknown_extension():
    with pytest.raises(InvalidFileTypeError):
        documents.extract_text("resume.exe", b"nope")


def test_extract_text_rejects_oversized_file():
    with pytest.raises(DocumentError, match="maximum allowed size"):
        documents.extract_text("resume.txt", b"x" * 100, max_size_bytes=10)


def test_extract_text_rejects_empty_document():
    with pytest.raises(EmptyDocumentError):
        documents.extract_text("resume.txt", b"   \n  ")


def test_extract_text_rejects_corrupt_docx():
    with pytest.raises(UnreadableDocumentError):
        documents.extract_text("resume.docx", b"not a real docx")


def test_extract_text_rejects_corrupt_pdf():
    with pytest.raises(UnreadableDocumentError):
        documents.extract_text("resume.pdf", b"not a real pdf")


def test_extract_text_is_case_insensitive_on_extension():
    text = documents.extract_text("RESUME.TXT", b"Upper case extension")
    assert text == "Upper case extension"


def test_docx_generation_produces_valid_document_with_sections():
    content = documents.build_tailored_resume_docx(RESUME_TEXT)

    assert content.startswith(b"PK")

    parsed = DocxDocument(BytesIO(content))
    paragraphs = [p.text for p in parsed.paragraphs if p.text.strip()]
    assert paragraphs[0] == "John Doe"
    assert "Professional Summary:" in paragraphs
    assert "Windows administration" in paragraphs
    assert any(
        p.runs and p.runs[0].bold
        for p in parsed.paragraphs
        if p.text.strip() == "Professional Summary:"
    )


def test_docx_generation_handles_empty_text():
    content = documents.build_tailored_resume_docx("")
    assert content.startswith(b"PK")
